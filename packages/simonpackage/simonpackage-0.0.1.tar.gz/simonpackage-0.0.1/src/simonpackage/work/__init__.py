# various functions for daily work routines


def getAttendanceList(name_list_str):
    """to get a name list from a name and occupation list of attendance,
    the name and occupation separated by at least two blank positions"""
    name_list_str = name_list_str.replace('\u3000', ' ')
    attendance = name_list_str.split('\n')
    names = []
    for i, a in enumerate(attendance):
        line = a.strip().split('  ')
        # print(f"line {i+1}: {line}")
        if len(line) >= 2:
            name = line[0].strip('（女）')
            names.append(name)
    return names

def insertTable(doc, l, rows, cols, priority='rows'):
    """
    insert a table populated with the elements of l into a docx object
    :param doc: a docx object
    :param l:
    :param rows:
    :param cols:
    :param priority: populate the table with row first or column first priority
    :return:
    """
    # columns are fixed, and rows could be not full
    table = doc.add_table(rows, cols)
    if priority == 'rows':
        # this is populating rows-first
        for i in range(rows):
            for j in range(cols):
                cell = table.cell(i, j)
                if i * cols + j < len(l):  # make sure the index wouldn't be out of range
                    cell.text = str(l[i * cols + j])
    elif priority == 'cols':
        # this is populating cols-first
        for i in range(cols):
            for j in range(rows):
                cell = table.cell(j, i)
                if i * rows + j < len(l):
                    cell.text = str(l[i * rows + j])
    else:
        raise "priority can only be set to 'rows' or 'cols'"


def convertTo2DList(l, rows, cols):
    """
     convert a list to a 2-d nested list with designated rows and cols
    :param l:  a list
    :param rows:
    :param cols:
    :return: a 2-d list which could have more elements than the original list
    """

    two_dim_list = []
    # fill the blank with None
    new_list = l[:]
    if len(names) < rows * cols:
        for i in range(rows * cols - len(l)):
            new_list.append(None)
    for i in range(len(new_list) // cols + 1):
        two_dim_list.append(new_list[i * cols:(i + 1) * cols])
    return two_dim_list


if __name__ == '__main__':
    from docx import Document
    doc = Document()
    insertTable(doc, names, 7, 9, 'cols')
    doc.save('a.docx')
